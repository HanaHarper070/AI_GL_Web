import json
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    # Load slides data
    with open('/Users/hanaharper/Downloads/AI_LG/slides_data.js', 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Strip 'const SLIDES_DATA = ' and trailing ';'
    json_str = content.replace('const SLIDES_DATA = ', '').strip()
    if json_str.endswith(';'):
        json_str = json_str[:-1]
    
    slides = json.loads(json_str)
    
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6] # Blank slide
    
    # Color palette
    NAVY_DARK = RGBColor(15, 23, 42)     # #0f172a
    NAVY_BLUE = RGBColor(30, 58, 138)    # #1e3a8a
    ROYAL_BLUE = RGBColor(29, 78, 216)   # #1d4ed8
    LIGHT_BG = RGBColor(248, 250, 252)   # #f8fafc
    WHITE = RGBColor(255, 255, 255)
    GRAY_TEXT = RGBColor(100, 116, 139)  # #64748b
    DARK_TEXT = RGBColor(30, 41, 59)     # #1e293b
    AMBER_GOLD = RGBColor(217, 119, 6)   # #d97706
    EMERALD_GREEN = RGBColor(16, 149, 106)
    CARD_BG = RGBColor(255, 255, 255)
    CARD_BORDER = RGBColor(226, 232, 240)

    for idx, s in enumerate(slides):
        slide = prs.slides.add_slide(blank_layout)
        
        # Add speaker notes
        if s.get('notes'):
            notes_slide = slide.notes_slide
            tf = notes_slide.notes_text_frame
            tf.text = s['notes']

        # Determine theme: Slide 1 and Slide 25 dark navy hero, others clean professional gov theme
        is_hero = s.get('type') == 'hero'
        
        if is_hero:
            # Dark Navy Background
            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
            bg.fill.solid()
            bg.fill.fore_color.rgb = NAVY_DARK
            bg.line.color.rgb = NAVY_DARK
            
            # Gold Accent Top Bar
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
            accent.fill.solid()
            accent.fill.fore_color.rgb = AMBER_GOLD
            accent.line.color.rgb = AMBER_GOLD
            
            # Badge
            badge_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(0.5))
            tf_b = badge_box.text_frame
            p_b = tf_b.paragraphs[0]
            p_b.text = s['badge'].upper()
            p_b.font.size = Pt(13)
            p_b.font.bold = True
            p_b.font.color.rgb = AMBER_GOLD
            
            # Title
            title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(11.3), Inches(1.2))
            tf_t = title_box.text_frame
            tf_t.word_wrap = True
            p_t = tf_t.paragraphs[0]
            p_t.text = s['title']
            p_t.font.size = Pt(30)
            p_t.font.bold = True
            p_t.font.color.rgb = WHITE
            
            # Subtitle
            sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(11.3), Inches(0.8))
            tf_s = sub_box.text_frame
            tf_s.word_wrap = True
            p_s = tf_s.paragraphs[0]
            p_s.text = s['subtitle']
            p_s.font.size = Pt(16)
            p_s.font.color.rgb = RGBColor(191, 219, 254)
            
            # 4 Cards Grid
            features = s['content'].get('features', [])
            card_w = Inches(2.6)
            card_h = Inches(2.5)
            gap = Inches(0.3)
            start_x = Inches(1.0)
            y = Inches(3.6)
            
            for i, f in enumerate(features[:4]):
                x = start_x + i * (card_w + gap)
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w, card_h)
                card.fill.solid()
                card.fill.fore_color.rgb = RGBColor(30, 41, 59)
                card.line.color.rgb = RGBColor(51, 65, 85)
                
                tf_c = card.text_frame
                tf_c.word_wrap = True
                tf_c.margin_left = Inches(0.2)
                tf_c.margin_right = Inches(0.2)
                tf_c.margin_top = Inches(0.3)
                
                p_ct = tf_c.paragraphs[0]
                p_ct.text = f['title']
                p_ct.font.size = Pt(15)
                p_ct.font.bold = True
                p_ct.font.color.rgb = WHITE
                
                p_cd = tf_c.add_paragraph()
                p_cd.space_before = Pt(10)
                p_cd.text = f['desc']
                p_cd.font.size = Pt(12)
                p_cd.font.color.rgb = RGBColor(148, 163, 184)

            # Footer
            foot_box = slide.shapes.add_textbox(Inches(1.0), Inches(6.6), Inches(11.3), Inches(0.4))
            tf_f = foot_box.text_frame
            p_f = tf_f.paragraphs[0]
            p_f.text = s['content'].get('footer_note', '')
            p_f.font.size = Pt(11)
            p_f.font.color.rgb = GRAY_TEXT

        else:
            # Professional Light Theme
            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
            bg.fill.solid()
            bg.fill.fore_color.rgb = LIGHT_BG
            bg.line.fill.background()

            # Top Header Bar
            header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.3))
            header.fill.solid()
            header.fill.fore_color.rgb = NAVY_BLUE
            header.line.fill.background()

            # Top Accent Line
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.25), Inches(13.333), Inches(0.05))
            accent.fill.solid()
            accent.fill.fore_color.rgb = AMBER_GOLD
            accent.line.fill.background()

            # Module Tag & Badge in Header
            tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(8.0), Inches(0.35))
            tf_tag = tag_box.text_frame
            p_tag = tf_tag.paragraphs[0]
            p_tag.text = f"{s['module'].upper()}  |  {s['badge']}"
            p_tag.font.size = Pt(11)
            p_tag.font.bold = True
            p_tag.font.color.rgb = AMBER_GOLD

            # Title in Header
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(10.5), Inches(0.7))
            tf_t = title_box.text_frame
            tf_t.word_wrap = True
            p_t = tf_t.paragraphs[0]
            p_t.text = s['title']
            p_t.font.size = Pt(20)
            p_t.font.bold = True
            p_t.font.color.rgb = WHITE

            # Slide Number in Header Right
            num_box = slide.shapes.add_textbox(Inches(11.5), Inches(0.4), Inches(1.2), Inches(0.5))
            tf_num = num_box.text_frame
            p_num = tf_num.paragraphs[0]
            p_num.alignment = PP_ALIGN.RIGHT
            p_num.text = f"{idx + 1} / {len(slides)}"
            p_num.font.size = Pt(13)
            p_num.font.bold = True
            p_num.font.color.rgb = RGBColor(191, 219, 254)

            # Subtitle below Header
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.4))
            tf_sub = sub_box.text_frame
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = s.get('subtitle', '')
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = GRAY_TEXT

            # Body based on type
            stype = s.get('type')
            
            if stype == 'prompt_display':
                # Large prompt card
                pcard = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.7), Inches(4.3))
                pcard.fill.solid()
                pcard.fill.fore_color.rgb = RGBColor(15, 23, 42)
                pcard.line.color.rgb = RGBColor(51, 65, 85)
                
                tf_p = pcard.text_frame
                tf_p.word_wrap = True
                tf_p.margin_left = Inches(0.4)
                tf_p.margin_right = Inches(0.4)
                tf_p.margin_top = Inches(0.3)
                
                p_p = tf_p.paragraphs[0]
                p_p.text = s['content']['prompt_text']
                p_p.font.size = Pt(12)
                p_p.font.color.rgb = RGBColor(226, 232, 240)

                # Key takeaway banner
                if s['content'].get('key_takeaway'):
                    tcard = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.7), Inches(0.65))
                    tcard.fill.solid()
                    tcard.fill.fore_color.rgb = RGBColor(239, 246, 255)
                    tcard.line.color.rgb = RGBColor(191, 219, 254)
                    
                    tf_tk = tcard.text_frame
                    p_tk = tf_tk.paragraphs[0]
                    p_tk.text = s['content']['key_takeaway']
                    p_tk.font.size = Pt(12)
                    p_tk.font.bold = True
                    p_tk.font.color.rgb = NAVY_BLUE

            elif stype in ['prompt_formula', 'two_column', 'comparison']:
                # Two or multi column layout
                if stype == 'prompt_formula':
                    elements = s['content']['elements']
                    for i, el in enumerate(elements):
                        row = i // 2
                        col = i % 2
                        x = Inches(0.8) + col * Inches(5.95)
                        y = Inches(1.9) + row * Inches(2.4)
                        
                        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.2))
                        card.fill.solid()
                        card.fill.fore_color.rgb = WHITE
                        card.line.color.rgb = CARD_BORDER
                        
                        tf_el = card.text_frame
                        tf_el.word_wrap = True
                        tf_el.margin_left = Inches(0.25)
                        tf_el.margin_right = Inches(0.25)
                        tf_el.margin_top = Inches(0.2)
                        
                        p_t = tf_el.paragraphs[0]
                        p_t.text = f"[{el['letter']}] {el['name']}"
                        p_t.font.size = Pt(14)
                        p_t.font.bold = True
                        p_t.font.color.rgb = NAVY_BLUE
                        
                        p_d = tf_el.add_paragraph()
                        p_d.space_before = Pt(4)
                        p_d.text = el['desc']
                        p_d.font.size = Pt(11)
                        p_d.font.color.rgb = DARK_TEXT
                        
                        p_ex = tf_el.add_paragraph()
                        p_ex.space_before = Pt(6)
                        p_ex.text = el['example']
                        p_ex.font.size = Pt(10)
                        p_ex.font.italic = True
                        p_ex.font.color.rgb = AMBER_GOLD

                elif stype == 'two_column':
                    # Left Card
                    card_l = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.75), Inches(5.0))
                    card_l.fill.solid()
                    card_l.fill.fore_color.rgb = WHITE
                    card_l.line.color.rgb = CARD_BORDER
                    
                    tf_l = card_l.text_frame
                    tf_l.word_wrap = True
                    tf_l.margin_left = Inches(0.3)
                    tf_l.margin_right = Inches(0.3)
                    tf_l.margin_top = Inches(0.3)
                    
                    p_lt = tf_l.paragraphs[0]
                    p_lt.text = s['content']['left_title']
                    p_lt.font.size = Pt(15)
                    p_lt.font.bold = True
                    p_lt.font.color.rgb = EMERALD_GREEN
                    
                    for item in s['content']['left_items']:
                        p_li = tf_l.add_paragraph()
                        p_li.space_before = Pt(8)
                        p_li.text = f"• {item}"
                        p_li.font.size = Pt(11)
                        p_li.font.color.rgb = DARK_TEXT

                    # Right Card
                    card_r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.9), Inches(5.75), Inches(5.0))
                    card_r.fill.solid()
                    card_r.fill.fore_color.rgb = WHITE
                    card_r.line.color.rgb = CARD_BORDER
                    
                    tf_r = card_r.text_frame
                    tf_r.word_wrap = True
                    tf_r.margin_left = Inches(0.3)
                    tf_r.margin_right = Inches(0.3)
                    tf_r.margin_top = Inches(0.3)
                    
                    p_rt = tf_r.paragraphs[0]
                    p_rt.text = s['content']['right_title']
                    p_rt.font.size = Pt(15)
                    p_rt.font.bold = True
                    p_rt.font.color.rgb = RGBColor(225, 29, 72)
                    
                    for item in s['content']['right_items']:
                        p_ri = tf_r.add_paragraph()
                        p_ri.space_before = Pt(8)
                        p_ri.text = f"• {item}"
                        p_ri.font.size = Pt(11)
                        p_ri.font.color.rgb = DARK_TEXT

                elif stype == 'comparison':
                    items = s['content']['items']
                    count = len(items)
                    cw = Inches(5.75) if count == 2 else Inches(3.75)
                    gap = Inches(0.2)
                    
                    for i, it in enumerate(items):
                        x = Inches(0.8) + i * (cw + gap)
                        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.9), cw, Inches(5.0))
                        card.fill.solid()
                        card.fill.fore_color.rgb = WHITE
                        card.line.color.rgb = CARD_BORDER
                        
                        tf_c = card.text_frame
                        tf_c.word_wrap = True
                        tf_c.margin_left = Inches(0.25)
                        tf_c.margin_right = Inches(0.25)
                        tf_c.margin_top = Inches(0.25)
                        
                        p_ct = tf_c.paragraphs[0]
                        p_ct.text = it['title']
                        p_ct.font.size = Pt(13)
                        p_ct.font.bold = True
                        p_ct.font.color.rgb = NAVY_BLUE
                        
                        p_cd = tf_c.add_paragraph()
                        p_cd.space_before = Pt(10)
                        p_cd.text = it['desc']
                        p_cd.font.size = Pt(11)
                        p_cd.font.color.rgb = DARK_TEXT

            elif stype == 'gov_types':
                types = s['content']['types']
                for i, t in enumerate(types):
                    row = i // 3
                    col = i % 3
                    x = Inches(0.8) + col * Inches(3.95)
                    y = Inches(1.9) + row * Inches(2.4)
                    
                    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.8), Inches(2.25))
                    card.fill.solid()
                    card.fill.fore_color.rgb = WHITE
                    card.line.color.rgb = CARD_BORDER
                    
                    tf_t = card.text_frame
                    tf_t.word_wrap = True
                    tf_t.margin_left = Inches(0.2)
                    tf_t.margin_right = Inches(0.2)
                    tf_t.margin_top = Inches(0.18)
                    
                    p_head = tf_t.paragraphs[0]
                    p_head.text = f"{t['num']}. {t['title']}"
                    p_head.font.size = Pt(13)
                    p_head.font.bold = True
                    p_head.font.color.rgb = NAVY_BLUE
                    
                    p_d = tf_t.add_paragraph()
                    p_d.space_before = Pt(4)
                    p_d.text = t['desc']
                    p_d.font.size = Pt(10.5)
                    p_d.font.color.rgb = DARK_TEXT
                    
                    p_det = tf_t.add_paragraph()
                    p_det.space_before = Pt(4)
                    p_det.text = t['detail']
                    p_det.font.size = Pt(9.5)
                    p_det.font.color.rgb = GRAY_TEXT

            elif stype == 'feature_card':
                features = s['content']['features']
                cw = Inches(3.75)
                gap = Inches(0.22)
                for i, f in enumerate(features):
                    x = Inches(0.8) + i * (cw + gap)
                    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.1), cw, Inches(4.7))
                    card.fill.solid()
                    card.fill.fore_color.rgb = WHITE
                    card.line.color.rgb = CARD_BORDER
                    
                    tf_f = card.text_frame
                    tf_f.word_wrap = True
                    tf_f.margin_left = Inches(0.3)
                    tf_f.margin_right = Inches(0.3)
                    tf_f.margin_top = Inches(0.3)
                    
                    p_ft = tf_f.paragraphs[0]
                    p_ft.text = f['title']
                    p_ft.font.size = Pt(15)
                    p_ft.font.bold = True
                    p_ft.font.color.rgb = NAVY_BLUE
                    
                    p_fd = tf_f.add_paragraph()
                    p_fd.space_before = Pt(12)
                    p_fd.text = f['desc']
                    p_fd.font.size = Pt(12)
                    p_fd.font.color.rgb = DARK_TEXT

            elif stype == 'grid_roadmap':
                modules = s['content']['modules']
                for i, m in enumerate(modules):
                    row = i // 4
                    col = i % 4
                    x = Inches(0.8) + col * Inches(2.95)
                    y = Inches(2.0) + row * Inches(2.4)
                    
                    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.8), Inches(2.2))
                    card.fill.solid()
                    card.fill.fore_color.rgb = WHITE
                    card.line.color.rgb = CARD_BORDER
                    
                    tf_m = card.text_frame
                    tf_m.word_wrap = True
                    tf_m.margin_left = Inches(0.2)
                    tf_m.margin_right = Inches(0.2)
                    tf_m.margin_top = Inches(0.2)
                    
                    p_c = tf_m.paragraphs[0]
                    p_c.text = m['code']
                    p_c.font.size = Pt(12)
                    p_c.font.bold = True
                    p_c.font.color.rgb = ROYAL_BLUE
                    
                    p_name = tf_m.add_paragraph()
                    p_name.space_before = Pt(4)
                    p_name.text = m['name']
                    p_name.font.size = Pt(12)
                    p_name.font.bold = True
                    p_name.font.color.rgb = NAVY_DARK
                    
                    p_md = tf_m.add_paragraph()
                    p_md.space_before = Pt(4)
                    p_md.text = m['desc']
                    p_md.font.size = Pt(10)
                    p_md.font.color.rgb = GRAY_TEXT

    out_path = '/Users/hanaharper/Downloads/AI_LG/AI_LG_Workshop_Presentation.pptx'
    prs.save(out_path)
    print(f"PowerPoint file saved successfully at: {out_path}")

if __name__ == '__main__':
    create_deck()
